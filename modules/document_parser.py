"""将支持的上传文件统一解析为纯文本。"""

import csv
from dataclasses import dataclass
from html.parser import HTMLParser
import io
import json
from pathlib import Path
import re


SUPPORTED_FORMATS = {".txt", ".md", ".csv", ".json", ".html", ".htm", ".pdf", ".docx", ".xlsx", ".pptx"}


@dataclass(frozen=True)
class ParsedDocument:
    """文档解析结果，供上传入库流程统一消费。"""

    name: str
    file_format: str
    text: str


class _TextOnlyHTMLParser(HTMLParser):
    """提取 HTML 可见文本，并跳过脚本和样式内容。"""

    # 初始化 HTML 可见文本收集器。
    def __init__(self) -> None:
        super().__init__(convert_charrefs=True)
        self.parts: list[str] = []
        self._ignored_depth = 0

    # 处理 HTML 开始标签并跳过脚本样式内容。
    def handle_starttag(self, tag: str, attrs: list[tuple[str, str | None]]) -> None:
        if tag.lower() in {"script", "style", "noscript"}:
            self._ignored_depth += 1
        elif tag.lower() in {"p", "div", "li", "br", "tr", "h1", "h2", "h3"}:
            self.parts.append("\n")

    # 处理 HTML 结束标签并维护忽略区域状态。
    def handle_endtag(self, tag: str) -> None:
        if tag.lower() in {"script", "style", "noscript"} and self._ignored_depth:
            self._ignored_depth -= 1
        elif tag.lower() in {"p", "div", "li", "tr", "h1", "h2", "h3"}:
            self.parts.append("\n")

    # 收集 HTML 中可见的文本内容。
    def handle_data(self, data: str) -> None:
        if not self._ignored_depth and data.strip():
            self.parts.append(data.strip())


# 统一换行并清理解析文本中的空白字符。
def _normalize(text: str) -> str:
    text = text.replace("\x00", " ").replace("\r\n", "\n").replace("\r", "\n")
    return re.sub(r"[ \t]+\n", "\n", re.sub(r"\n{3,}", "\n\n", text)).strip()


# 按文件格式解析内容并返回统一文档结构。
def parse_document(filename: str, content: bytes) -> ParsedDocument:
    """按扩展名解析文档；不支持的格式或空文本会明确报错。"""
    suffix = Path(filename).suffix.lower()
    if suffix not in SUPPORTED_FORMATS:
        raise ValueError(f"暂不支持 {suffix or '无扩展名'} 格式")
    try:
        if suffix in {".txt", ".md"}:
            text = content.decode("utf-8-sig")
        elif suffix == ".csv":
            rows = csv.reader(io.StringIO(content.decode("utf-8-sig")))
            text = "\n".join(" | ".join(cell.strip() for cell in row) for row in rows)
        elif suffix == ".json":
            text = json.dumps(json.loads(content.decode("utf-8-sig")), ensure_ascii=False, indent=2)
        elif suffix in {".html", ".htm"}:
            parser = _TextOnlyHTMLParser()
            parser.feed(content.decode("utf-8-sig"))
            text = "\n".join(parser.parts)
        elif suffix == ".pdf":
            from pypdf import PdfReader
            reader = PdfReader(io.BytesIO(content))
            text = "\n".join(page.extract_text() or "" for page in reader.pages)
        elif suffix == ".docx":
            from docx import Document
            document = Document(io.BytesIO(content))
            text = "\n".join(paragraph.text for paragraph in document.paragraphs)
            for table in document.tables:
                text += "\n" + "\n".join(" | ".join(cell.text for cell in row.cells) for row in table.rows)
        elif suffix == ".xlsx":
            from openpyxl import load_workbook
            workbook = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
            text = "\n".join(
                f"工作表：{sheet.title}\n" + "\n".join(
                    " | ".join("" if value is None else str(value) for value in row)
                    for row in sheet.iter_rows(values_only=True)
                )
                for sheet in workbook.worksheets
            )
            workbook.close()
        else:
            from pptx import Presentation
            presentation = Presentation(io.BytesIO(content))
            text = "\n".join(
                shape.text for slide in presentation.slides for shape in slide.shapes if shape.has_text_frame
            )
    except (UnicodeDecodeError, json.JSONDecodeError, OSError, KeyError, ValueError) as error:
        raise ValueError("文件内容无法解析，请确认文件未损坏且文本编码为 UTF-8") from error

    normalized = _normalize(text)
    if not normalized:
        raise ValueError("文档中没有可检索的文本内容")
    return ParsedDocument(name=Path(filename).name, file_format=suffix.lstrip("."), text=normalized)
